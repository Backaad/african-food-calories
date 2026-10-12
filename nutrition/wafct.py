#!/usr/bin/env python3
"""Fetch the FAO/INFOODS West African Food Composition Table (WAFCT 2019) and query it.

    python nutrition/wafct.py fetch                 # find and download the Excel file
    python nutrition/wafct.py describe              # sheet names and header rows
    python nutrition/wafct.py search cassava palm   # foods whose name matches, with key nutrients

The Excel file is not committed (see nutrition/README.md for its licence); only the
rows used by the recipes are extracted, with their WAFCT codes.
"""

import argparse
import re
import sys
from pathlib import Path

NUTRITION_DIR = Path(__file__).resolve().parent
CACHE = NUTRITION_DIR / "cache"
XLSX = CACHE / "WAFCT_2019.xlsx"
USER_AGENT = "african-food-calories-nutrition/0.1 (+https://github.com/backaad/african-food-calories)"

# Pages that link to the WAFCT 2019 Excel datasheets.
LANDING_PAGES = [
    "https://www.fao.org/infoods/infoods/tables-and-databases/faoinfoods-databases/en/",
    "https://www.fao.org/infoods/infoods/tables-and-databases/africa/en/",
]

# Components we use, matched against the English header row of the
# "NV_sum" sheets (INFOODS tagname in brackets).
COMPONENTS = {
    "energy_kcal": r"^Energy \(kcal\)",  # ENERC
    "water_g": r"^Water \(g\)",  # WATER
    "protein_g": r"^Protein, total",  # PROTCNT
    "fat_g": r"^Fat, total",  # FAT
    "carbohydrate_g": r"^Carbohydrate, available",  # CHOAVLDF
    "fibre_g": r"^Fibre, total dietary",  # FIBTG
}
MAIN_SHEET = "03 NV_sum_39 (per 100g EP)"


def find_excel_urls(html: str, base: str) -> list[str]:
    from urllib.parse import urljoin

    links = re.findall(r'href="([^"]+?\.xlsx?)"', html, flags=re.I)
    urls = [urljoin(base, link) for link in links]
    return sorted({u for u in urls if re.search(r"wafct|west|western|afrique", u, re.I)}) or sorted(set(urls))


def fetch() -> Path:
    import requests

    CACHE.mkdir(exist_ok=True)
    if XLSX.exists():
        return XLSX
    session = requests.Session()
    session.headers["User-Agent"] = USER_AGENT
    candidates = []
    for page in LANDING_PAGES:
        try:
            response = session.get(page, timeout=60)
            response.raise_for_status()
        except Exception as exc:
            print(f"  {page}: {exc}", file=sys.stderr)
            continue
        found = find_excel_urls(response.text, page)
        print(f"{page}: {found}", file=sys.stderr)
        candidates += found
    for url in candidates:
        response = session.get(url, timeout=120)
        if response.ok and response.content[:2] == b"PK":  # xlsx files are zip archives
            XLSX.write_bytes(response.content)
            print(f"downloaded {url} ({len(response.content)} bytes)", file=sys.stderr)
            (CACHE / "SOURCE_URL").write_text(url + "\n")
            return XLSX
        print(f"  not an xlsx: {url} ({response.status_code})", file=sys.stderr)
    sys.exit("No WAFCT Excel file found on the landing pages.")


def _norm(text) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def find_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    """Locate the English header row and the column of each field we need."""
    for i, row in enumerate(rows[:30]):
        cells = [_norm(c) for c in row]
        cols = {}
        for j, cell in enumerate(cells):
            low = cell.lower()
            if low == "code":
                cols.setdefault("code", j)
            elif low == "food name in english":
                cols.setdefault("name_en", j)
            elif low == "food name in french":
                cols.setdefault("name_fr", j)
            for key, pattern in COMPONENTS.items():
                if re.search(pattern, cell):
                    cols.setdefault(key, j)
        if {"code", "name_en", *COMPONENTS} <= set(cols):
            return i, cols
    raise ValueError("header row not found")


def load_table(path: Path = XLSX, sheet: str = MAIN_SHEET) -> tuple[str, list[dict]]:
    """Return (sheet name, foods as dicts) from the main composition sheet."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    header_index, cols = find_header(rows)
    records = []
    for row in rows[header_index + 1:]:
        code = _norm(row[cols["code"]]) if cols["code"] < len(row) else ""
        if not re.match(r"^\d+_\d+$", code):
            continue  # food group title rows, tagname row, notes
        records.append({"code": code} | {k: row[j] for k, j in cols.items() if k != "code" and j < len(row)})
    if not records:
        raise ValueError(f"no foods found in sheet {sheet!r} of {path}")
    return sheet, records


def to_number(value) -> float | None:
    """WAFCT writes estimated values in brackets, e.g. [1.2], and traces as 'tr'."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().strip("[]()").replace(",", ".")
    if text.lower() in {"tr", "trace"}:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return None


def describe(path: Path = XLSX) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines = [f"sheets: {wb.sheetnames}"]
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(min_row=1, max_row=6, values_only=True))
        lines.append(f"\n## {name}")
        for row in rows:
            lines.append(" | ".join(_norm(c)[:40] for c in row[:25]))
    return "\n".join(lines)


def search(terms: list[str], path: Path = XLSX, limit: int = 15) -> str:
    sheet, records = load_table(path)
    lines = [f"sheet: {sheet}, {len(records)} foods"]
    for term in terms:
        pattern = re.compile(term, re.I)
        hits = [r for r in records if pattern.search(_norm(r.get("name_en"))) or pattern.search(_norm(r.get("name_fr")))]
        lines.append(f"\n## {term} ({len(hits)} hits)")
        for r in hits[:limit]:
            values = " ".join(f"{k.split('_')[0][:4]}={_norm(r.get(k))}" for k in COMPONENTS)
            lines.append(f"{r['code']} | {_norm(r.get('name_en'))[:90]} | {_norm(r.get('name_fr'))[:60]} | {values}")
    return "\n".join(lines)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("fetch")
    sub.add_parser("describe")
    s = sub.add_parser("search")
    s.add_argument("terms", nargs="*")
    s.add_argument("--file", type=Path, help="one search pattern per line")
    s.add_argument("--limit", type=int, default=15)
    args = parser.parse_args(argv)

    if args.command == "fetch":
        print(fetch())
    elif args.command == "describe":
        print(describe(fetch()))
    elif args.command == "search":
        terms = list(args.terms)
        if args.file:
            terms += [t.strip() for t in args.file.read_text(encoding="utf-8").splitlines() if t.strip()]
        print(search(terms, fetch(), args.limit))


if __name__ == "__main__":
    main()
